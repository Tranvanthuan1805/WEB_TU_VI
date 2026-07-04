import openpyxl
import pymongo
import datetime
import sys

# Ensure output displays Vietnamese characters correctly
sys.stdout.reconfigure(encoding='utf-8')

excel_path = r"c:\Users\duotech\OneDrive\Desktop\DU_AN_WEB_TU_VI\Project-2026\BattuHL-Tubinh-Tuvi-Quedich-LTT\BattuHL-Tubinh-Tuvi-Quedich-LTT.xlsx"
mongo_uri = "mongodb+srv://duotechcompanyhr_db_user:P3GRlMw7enKdlwU2@cluster0.y0urkmd.mongodb.net/?appName=Cluster0"
db_name = "tuvi_db"

def parse_date(val):
    if not val:
        return None
    if isinstance(val, datetime.datetime):
        return val
    if isinstance(val, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"):
            try:
                return datetime.datetime.strptime(val.strip(), fmt)
            except ValueError:
                continue
    return None

def main():
    print(f"Connecting to MongoDB at {mongo_uri}...")
    client = pymongo.MongoClient(mongo_uri)
    db = client[db_name]

    print(f"Loading Excel workbook: {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # 1. Process '384Hào' Sheet
    print("Processing '384Hào' sheet...")
    sheet_384 = wb['384Hào']
    hexagrams_col = db['hexagram_lines']
    hexagrams_col.drop()  # Clear existing

    hexagram_docs = []
    # Columns 3 to 66 represent the 64 hexagrams
    for col in range(3, 67):
        index_val = sheet_384.cell(2, col).value
        if index_val is None:
            continue
        try:
            index = int(index_val)
        except ValueError:
            continue

        name = sheet_384.cell(3, col).value
        description = sheet_384.cell(4, col).value
        
        # General poems (Rows 5 to 8)
        general_poems = []
        for row in range(5, 9):
            val = sheet_384.cell(row, col).value
            if val is not None and str(val).strip():
                general_poems.append(str(val).strip())

        # 6 lines (Hào 6 down to Hào 1, starting at row 9)
        lines = []
        for line_num in range(1, 7):
            # Hào 6 is at row 9, Hào 5 at 20, Hào 4 at 31, etc.
            row_start = 9 + (6 - line_num) * 11
            
            meaning = sheet_384.cell(row_start, col).value
            interpretation = sheet_384.cell(row_start + 1, col).value
            mhc = sheet_384.cell(row_start + 2, col).value
            mkh = sheet_384.cell(row_start + 3, col).value
            qc = sheet_384.cell(row_start + 4, col).value
            gs = sheet_384.cell(row_start + 5, col).value
            nt = sheet_384.cell(row_start + 6, col).value

            # Line poems (row_start + 7 to row_start + 10)
            line_poems = []
            for row in range(row_start + 7, row_start + 11):
                val = sheet_384.cell(row, col).value
                if val is not None and str(val).strip():
                    line_poems.append(str(val).strip())

            lines.append({
                "line_number": line_num,
                "meaning": str(meaning).strip() if meaning is not None else None,
                "interpretation": str(interpretation).strip() if interpretation is not None else None,
                "mhc": str(mhc).strip() if mhc is not None else None,
                "mkh": str(mkh).strip() if mkh is not None else None,
                "qc": str(qc).strip() if qc is not None else None,
                "gs": str(gs).strip() if gs is not None else None,
                "nt": str(nt).strip() if nt is not None else None,
                "poems": line_poems
            })

        doc = {
            "index": index,
            "name": str(name).strip() if name is not None else None,
            "description": str(description).strip() if description is not None else None,
            "general_poems": general_poems,
            "lines": lines
        }
        hexagram_docs.append(doc)

    if hexagram_docs:
        hexagrams_col.insert_many(hexagram_docs)
    print(f"Imported {len(hexagram_docs)} hexagram lines to 'hexagram_lines' collection.")

    # 2. Process '64Quẻ' Sheet
    print("Processing '64Quẻ' sheet...")
    sheet_64 = wb['64Quẻ']
    details_col = db['hexagram_details']
    details_col.drop()

    detail_docs = []
    # Columns are in pairs starting from 2 to 141 (c is base, c+1 is alternate)
    last_binary_code = None
    for col in range(2, 142, 2):
        name = sheet_64.cell(5, col).value
        if not name:
            continue
        
        # Read binary code from row 4 of base column, or fallback to last seen
        bin_val = sheet_64.cell(4, col).value
        if bin_val is not None:
            last_binary_code = str(bin_val).strip()
        binary_code = last_binary_code

        # Parse 6 lines (Row 11 is line 1, Row 6 is line 6)
        lines = []
        for line_num in range(1, 7):
            row_idx = 12 - line_num
            yang_yin = sheet_64.cell(row_idx, col).value
            stem_branch = sheet_64.cell(row_idx, col + 1).value

            lines.append({
                "line_number": line_num,
                "yang_yin": str(yang_yin).strip() if yang_yin is not None else None,
                "stem_branch": str(stem_branch).strip() if stem_branch is not None else None
            })

        # Yang branches (rows 12 to 17)
        duong_branches = []
        for row in range(12, 18):
            val = sheet_64.cell(row, col).value
            duong_branches.append(str(val).strip() if val is not None else None)

        # Yin branches (rows 18 to 23)
        am_branches_col = []
        am_branches_alt = []
        for row in range(18, 24):
            val1 = sheet_64.cell(row, col).value
            val2 = sheet_64.cell(row, col + 1).value
            am_branches_col.append(str(val1).strip() if val1 is not None else None)
            am_branches_alt.append(str(val2).strip() if val2 is not None else None)

        rating = sheet_64.cell(24, col).value
        dac_thoi = sheet_64.cell(25, col).value
        extra_val = sheet_64.cell(26, col).value or sheet_64.cell(26, col + 1).value

        doc = {
            "name": str(name).strip(),
            "binary_code": binary_code,
            "lines": lines,
            "duong_branches": duong_branches,
            "am_branches": am_branches_col,
            "am_branches_alt": am_branches_alt,
            "rating": str(rating).strip() if rating is not None else None,
            "dac_thoi": int(dac_thoi) if dac_thoi is not None else None,
            "extra_value": int(extra_val) if extra_val is not None else None
        }
        detail_docs.append(doc)

    if detail_docs:
        details_col.insert_many(detail_docs)
    print(f"Imported {len(detail_docs)} details to 'hexagram_details' collection.")

    # 3. Process 'TiếtKhí' Sheet
    print("Processing 'TiếtKhí' sheet...")
    sheet_tietkhi = wb['TiếtKhí']
    tietkhi_col = db['solar_terms']
    tietkhi_col.drop()

    tietkhi_docs = []
    # Rows 3 onwards
    for r in range(3, sheet_tietkhi.max_row + 1):
        code = sheet_tietkhi.cell(r, 2).value
        term_name = sheet_tietkhi.cell(r, 3).value
        dt_val = sheet_tietkhi.cell(r, 4).value

        if code is None or term_name is None:
            continue

        entry_time = parse_date(dt_val)

        doc = {
            "code": str(code).strip(),
            "term_name": str(term_name).strip(),
            "entry_time": entry_time
        }
        tietkhi_docs.append(doc)

    if tietkhi_docs:
        tietkhi_col.insert_many(tietkhi_docs)
    print(f"Imported {len(tietkhi_docs)} solar terms to 'solar_terms' collection.")

    # 4. Process 'Nhập liệu' Sheet (User Profiles)
    print("Processing 'Nhập liệu' sheet...")
    sheet_nhaplieu = wb['Nhập liệu']
    profiles_col = db['profiles']
    profiles_col.drop()

    profile_docs = []
    # Rows 6 to 1004
    for r in range(6, 1005):
        name = sheet_nhaplieu.cell(r, 3).value
        if not name or not str(name).strip():
            continue

        tt = sheet_nhaplieu.cell(r, 2).value
        gender = sheet_nhaplieu.cell(r, 4).value
        day = sheet_nhaplieu.cell(r, 5).value
        month = sheet_nhaplieu.cell(r, 6).value
        year = sheet_nhaplieu.cell(r, 7).value
        hour = sheet_nhaplieu.cell(r, 8).value
        minute = sheet_nhaplieu.cell(r, 9).value
        notes = sheet_nhaplieu.cell(r, 10).value
        dt_val = sheet_nhaplieu.cell(r, 12).value

        birth_datetime = parse_date(dt_val)

        doc = {
            "index": int(tt) if tt is not None else None,
            "fullname": str(name).strip(),
            "gender": str(gender).strip() if gender is not None else None,
            "birth_day": int(day) if day is not None else None,
            "birth_month": int(month) if month is not None else None,
            "birth_year": int(year) if year is not None else None,
            "birth_hour": int(hour) if hour is not None else None,
            "birth_minute": int(minute) if minute is not None else None,
            "notes": str(notes).strip() if notes is not None else None,
            "birth_datetime": birth_datetime
        }
        profile_docs.append(doc)

    if profile_docs:
        profiles_col.insert_many(profile_docs)
    print(f"Imported {len(profile_docs)} profiles to 'profiles' collection.")

    print("\nAll database tables successfully migrated to MongoDB!")
    
    # Simple count verification printout
    print("\n--- Summary Verification ---")
    print(f"hexagram_lines: {db['hexagram_lines'].count_documents({})}")
    print(f"hexagram_details: {db['hexagram_details'].count_documents({})}")
    print(f"solar_terms: {db['solar_terms'].count_documents({})}")
    print(f"profiles: {db['profiles'].count_documents({})}")

if __name__ == "__main__":
    main()
